"""Arma proveedores.xlsx: una hoja por rubro, ordenada por prioridad, con columnas para anotar.

Uso:  pip install openpyxl
      python3 proveedores/hacer_excel.py
Datos de la búsqueda del 23/10/2026 (directorios y buscadores: confirmar por teléfono).
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# (prioridad, negocio, qué pedirle, dirección, teléfono, otro contacto, fuente)
HOJAS = {
    "Madera": (
        "Pedir: MDF/fibrofácil CRUDO 3 mm y terciado 3 mm (no fenólico). Precio por placa y "
        "cuánto cobran por cortarla en hojas de 60 x 40 cm.",
        [
            (1, "Placacentro Distribuidora Lamadrid (Roca)", "MDF 3 mm, corte a medida, envío",
             "Av. Roca 4000", "0381 436-7500", "L a V 8-12 y 15:30-19:30, sáb 8:30-12:30",
             "Web Placacentro"),
            (2, "Placacentro Distribuidora Lamadrid (central)", "MDF 3 mm, corte a medida",
             "Av. Ejército del Norte 395", "0381 433-0635", "", "Web Placacentro"),
            (3, "Macromad Maderas", "Placas Faplac, cortes en el día", "Av. Belgrano 1965",
             "381 319-0352", "Facebook: macromadok", "Directorio + Facebook"),
            (4, "Maderera San Carlos SRL", "MDF y TERCIADO", "Av. Ejército del Norte 1011",
             "0381 433-1143", "", "Directorio"),
            (5, "Madernoa", "Maderera grande", "Av. Juan B. Justo 1527", "381 428-3250",
             "Otro: 0381 428-2087", "Directorio + Facebook"),
            (6, "MS Maderas", "MDF y derivados", "Asunción 1464", "381 456-2095",
             "Facebook: msmaderastuc", "Facebook"),
            (7, "Maderplak", "Placas (sin confirmar)", "Av. Gral. Roca 1999", "381 452-5930",
             "Facebook: maderplak", "Directorio"),
            (8, "Ternor SRL", "Fibrofácil en hojas chicas (artesanías)", "Marcos Paz 1973",
             "381 432-4260", "Otro: 0381 423-1207", "Directorio"),
            (9, "Hipermaderas SRL", "Maderas", "Av. Ejército del Norte 640", "0381 432-3439", "",
             "Directorio viejo"),
            (10, "maderTUC", "Maderera", "Lavalle 3655", "381 614-6950", "Celular (WhatsApp)",
             "Directorio"),
            (11, "Maderas Jujuy", "Maderera", "Pcia. de Jujuy 2906", "381 256-5633",
             "Otro: 381 638-9330", "Directorio"),
            (12, "Quebrachito Maderera", "Maderera", "Calle Génova s/n, Tafí Viejo", "381 343-0701",
             "", "Directorio"),
        ]),
    "Electrónica": (
        "Pedir de a 10 a 50: motor 130 de 3 V con cables, portapilas 2xAA con cables (ideal con "
        "tapa a tornillo), pulsador 12 mm CON TRABA, bornes tipo Wago 221 de 2 vías, cable fino "
        "rojo y negro.",
        [
            (1, "Electrónica Santa Fe", "Componentes", "Santa Fe 2020", "381 524-7386",
             "WhatsApp. Fijo: 0381 423-5381", "Web propia"),
            (2, "Computer-C", "Componentes, pilas, cables", "José Colombres 247", "381 333-7744",
             "Otro: 381 494-8888 · ventas@computer-c.com.ar", "Web propia"),
            (3, "ElectroGeek", "Arduino, robótica, motores", "Combate de las Piedras 388",
             "381 657-0242", "IG: @electrogeekarg", "Directorio"),
            (4, "Genaro Electronics", "Repuestos de electrónica", "9 de Julio 618", "0381 424-3951",
             "", "Directorio"),
            (5, "Electrónica Centro", "Service (preguntar repuestos)", "J. Colombres 390",
             "0381 421-5763", "", "Directorio"),
            (6, "Electrónica Ariez", "Componentes", "Blvd. 9 de Julio 398, Yerba Buena",
             "0381 435-1226", "", "Directorio"),
            (7, "Electrónica Yerba Buena", "Electrónica", "R. de Escalada 146, Yerba Buena",
             "0381 425-3413", "", "Directorio"),
            (8, "Santiago Kohn S.A.", "Bornes Wago y cable", "Maipú 471 / Av. Aconquija 755 (YB)",
             "no encontrado", "santiagokohn.com.ar", "Web"),
            (9, "TodoMicro (Buenos Aires, envía)", "Mayorista: motores y portapilas",
             "J. B. Justo 3046, CABA", "+54 9 11 2573-0643", "WhatsApp mayorista", "Web propia"),
            (10, "SYC Semiconductores (Buenos Aires, envía)", "Mayor y menor", "Paraná 274, CABA",
             "011 4371-5300", "+54 9 11 2590-3240", "Web propia"),
            (11, "Microelectrónica Componentes (Bs. As.)", "Componentes", "Paraná 180, CABA",
             "011 4375-0066", "WhatsApp +54 9 11 6273-4869", "Web propia"),
            (12, "Electrocomponentes S.A. (Bs. As., envía)", "Distribuidor grande", "Solís 225, CABA",
             "011 4375-3366", "WhatsApp +54 9 11 5379-1151", "Directorio"),
        ]),
    "O-rings (cubiertas)": (
        "Pedir: o-ring de nitrilo (NBR) 31 x 3 mm (interior x cordón), o 30 x 3, o 32 x 3. "
        "Precio por 200 unidades (4 por auto).",
        [
            (1, "Rodaever S.A.", "Rulemanes, retenes", "José Colombres 340", "0381 422-3626", "",
             "Directorio"),
            (2, "Rulemanes 488", "Retenes (envía a todo el país)", "J. Colombres 488",
             "0381 400-5415", "IG: @rulemanes488tucuman", "Directorio + Instagram"),
            (3, "Centro Rodamientos SRL", "Distribuidor SKF", "San Juan 1040", "381 421-2720",
             "Facebook", "Directorio"),
            (4, "Movimec S.A.", "Retenes y rodamientos", "Catamarca 416 (o Av. Alem 333)",
             "0381 424-2513", "Otro: 0381 424-2514", "Directorio"),
            (5, "Rulemanes Salta SRL", "Retenes", "Catamarca 460", "381 431-1952", "", "Directorio"),
            (6, "Rulemanes San Martín SRL", "Retenes", "José Colombres 110", "381 642-5300", "",
             "Directorio"),
            (7, "Tucumán Orugas", "Dice tener o-rings", "Laprida 2780", "no encontrado",
             "tucumanorugas.com.ar", "Web"),
        ]),
    "Otros": (
        "Ruedas ya hechas (opcional: las de madera ya vienen en las hojas de corte) y "
        "talleres de corte láser de referencia.",
        [
            (1, "Tienda de Arte (librería artística)", "Preguntar por ruedas de maqueta",
             "San Juan 955", "381 409-8044", "Yerba Buena: Av. Aconquija 1273, 381 510-6349",
             "Web propia"),
            (2, "ElectroGeek", "Ruedas de robótica", "Combate de las Piedras 388", "381 657-0242",
             "", "Directorio"),
            (3, "Tucumán Artística MDF", "Corte láser / fibrofácil", "Barrio 260 Viviendas",
             "0381 682-8767", "Facebook", "Directorio"),
            (4, "Laser Guimi", "Corte láser", "Sáenz Peña 775", "no encontrado", "", "Directorio"),
            (5, "Noagraf", "Corte láser", "Pasaje Padilla 29", "no encontrado", "", "Directorio"),
        ]),
}

COLUMNAS = ["#", "Negocio", "Qué tiene", "Dirección", "Teléfono", "Otro contacto / horario",
            "Fuente del dato", "¿Llamé?", "Precio que me dieron", "Notas"]
ANCHOS = [4, 38, 30, 32, 18, 34, 18, 9, 22, 30]


def main():
    wb = Workbook()
    wb.remove(wb.active)
    titulo = Font(bold=True, size=13, color="FFFFFF")
    cab = PatternFill("solid", fgColor="8A4B12")
    nota = PatternFill("solid", fgColor="FFF4E0")
    anotar = PatternFill("solid", fgColor="FFFBEA")
    for nombre, (pedido, filas) in HOJAS.items():
        ws = wb.create_sheet(nombre)
        ws["A1"] = pedido
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNAS))
        ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
        ws["A1"].fill = nota
        ws["A1"].font = Font(bold=True, size=11)
        ws.row_dimensions[1].height = 36
        ws.append(COLUMNAS)
        for c in ws[2]:
            c.font = titulo
            c.fill = cab
            c.alignment = Alignment(vertical="center")
        for fila in sorted(filas):
            ws.append(list(fila) + ["", "", ""])
        for fila in ws.iter_rows(min_row=3):
            for c in fila:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            fila[4].font = Font(bold=True, size=12)
            for c in fila[7:]:
                c.fill = anotar
        for i, w in enumerate(ANCHOS, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "B3"
    ws = wb.create_sheet("Leeme")
    for linea in [
        "Proveedores para el kit de autos - búsqueda del 23/10/2026",
        "Los datos salen de buscadores y directorios: pueden estar desactualizados. Confirmar al llamar.",
        "Cada hoja está ordenada por prioridad (#1 = llamar primero).",
        "Las columnas amarillas son para anotar: ¿Llamé?, precio y notas.",
        "Precio de referencia: placa MDF 3 mm 1,83 x 2,60 m entre $16.000 y $27.000.",
        "Electrónica por auto: $4.500 a $6.000 comprando de a 10 (Mercado Libre).",
    ]:
        ws.append([linea])
    ws.column_dimensions["A"].width = 100
    ws["A1"].font = Font(bold=True, size=13)
    ruta = os.path.join(os.path.dirname(os.path.abspath(__file__)), "proveedores.xlsx")
    wb.save(ruta)
    print(ruta)


if __name__ == "__main__":
    main()
